from fastapi import FastAPI, File, UploadFile, Form
from fastapi.middleware.cors import CORSMiddleware
from groq import Groq
from sklearn.feature_extraction.text import TfidfVectorizer
from sklearn.metrics.pairwise import cosine_similarity
import pdfplumber
import docx
import io
import os
import json
from typing import List
from dotenv import load_dotenv

load_dotenv()

app = FastAPI(title="Job Suite AI API", version="1.0.0")

app.add_middleware(
    CORSMiddleware,
    allow_origins=["*"],
    allow_credentials=True,
    allow_methods=["*"],
    allow_headers=["*"],
)

client = Groq(api_key=os.environ.get("GROQ_API_KEY"))


# ── Helpers ────────────────────────────────────────────────────────────────────

def extract_text_from_pdf(file_bytes: bytes) -> str:
    text = ""
    with pdfplumber.open(io.BytesIO(file_bytes)) as pdf:
        for page in pdf.pages:
            page_text = page.extract_text()
            if page_text:
                text += page_text + "\n"
    return text.strip()


def extract_text_from_docx(file_bytes: bytes) -> str:
    doc = docx.Document(io.BytesIO(file_bytes))
    return "\n".join([para.text for para in doc.paragraphs]).strip()


def extract_text(filename: str, file_bytes: bytes) -> str:
    if filename.endswith(".pdf"):
        return extract_text_from_pdf(file_bytes)
    elif filename.endswith(".docx"):
        return extract_text_from_docx(file_bytes)
    return ""


def extract_skills(text: str) -> List[str]:
    common_skills = [
        "python", "javascript", "typescript", "react", "next.js", "node.js",
        "fastapi", "django", "flask", "postgresql", "mysql", "mongodb", "redis",
        "docker", "kubernetes", "aws", "gcp", "azure", "git", "ci/cd",
        "machine learning", "deep learning", "nlp", "sql", "rest api",
        "graphql", "tailwind", "html", "css", "java", "c++", "go", "rust",
        "scikit-learn", "tensorflow", "pytorch", "leadership", "communication",
        "agile", "scrum", "figma", "ui/ux",
    ]
    text_lower = text.lower()
    return [skill for skill in common_skills if skill in text_lower]


def ask_groq(prompt: str) -> str:
    response = client.chat.completions.create(
        model="llama-3.3-70b-versatile",
        messages=[{"role": "user", "content": prompt}],
        temperature=0.3,
    )
    return response.choices[0].message.content.strip()


def parse_json_response(raw: str) -> dict:
    if raw.startswith("```"):
        parts = raw.split("```")
        raw = parts[1] if len(parts) > 1 else raw
        if raw.startswith("json"):
            raw = raw[4:]
    return json.loads(raw.strip())


# ── Routes ─────────────────────────────────────────────────────────────────────

@app.get("/")
def root():
    return {"status": "Job Suite AI API is running"}


@app.post("/analyze-resume")
async def analyze_resume(resume: UploadFile = File(...)):
    file_bytes = await resume.read()
    text = extract_text(resume.filename, file_bytes)

    if not text:
        return {"error": "Could not extract text from resume"}

    prompt = f"""You are an expert resume reviewer with 10+ years of hiring experience.

Analyze this resume and respond in this exact JSON format with no extra text or markdown:
{{
  "score": <number 0-100>,
  "summary": "<2 sentence overall assessment>",
  "strengths": ["<strength 1>", "<strength 2>", "<strength 3>"],
  "weaknesses": ["<weakness 1>", "<weakness 2>", "<weakness 3>"],
  "improvements": ["<specific improvement 1>", "<specific improvement 2>", "<specific improvement 3>"],
  "keywords_missing": ["<keyword 1>", "<keyword 2>", "<keyword 3>"]
}}

Resume:
{text[:3000]}"""

    raw = ask_groq(prompt)
    result = parse_json_response(raw)
    return result


@app.post("/generate-cover-letter")
async def generate_cover_letter(
    resume: UploadFile = File(...),
    job_description: str = Form(...),
    tone: str = Form(default="professional"),
):
    file_bytes = await resume.read()
    resume_text = extract_text(resume.filename, file_bytes)

    if not resume_text:
        return {"error": "Could not extract text from resume"}

    prompt = f"""You are an expert career coach who writes outstanding cover letters.

Write a cover letter with a {tone} tone based on this resume and job description.

Requirements:
- 3-4 paragraphs
- Tailored specifically to this job
- Highlight relevant experience from the resume
- Do not use generic phrases like "I am writing to apply"
- End with a confident call to action
- No placeholders — write it as if ready to send

Job Description:
{job_description[:2000]}

Resume:
{resume_text[:2000]}

Respond with ONLY the cover letter text, nothing else."""

    cover_letter = ask_groq(prompt)
    return {"cover_letter": cover_letter}


@app.post("/match-job")
async def match_job(
    resume: UploadFile = File(...),
    job_description: str = Form(...),
):
    file_bytes = await resume.read()
    resume_text = extract_text(resume.filename, file_bytes)

    if not resume_text:
        return {"error": "Could not extract text from resume"}

    # TF-IDF match score
    vectorizer = TfidfVectorizer(stop_words="english")
    tfidf = vectorizer.fit_transform([job_description, resume_text])
    score = float(cosine_similarity(tfidf[0], tfidf[1])[0][0])

    # Skill gap analysis
    resume_skills = extract_skills(resume_text)
    jd_skills = extract_skills(job_description)
    missing = [s for s in jd_skills if s not in resume_skills]
    matched = [s for s in jd_skills if s in resume_skills]

    prompt = f"""You are a career advisor. Based on this resume and job description, give 3 specific actionable tips to improve the candidate's chances.

Respond in this exact JSON format with no extra text or markdown:
{{
  "tips": ["<tip 1>", "<tip 2>", "<tip 3>"]
}}

Job Description:
{job_description[:1500]}

Resume:
{resume_text[:1500]}"""

    raw = ask_groq(prompt)
    ai_result = parse_json_response(raw)

    return {
        "match_score": round(score * 100, 1),
        "matched_skills": matched,
        "missing_skills": missing,
        "tips": ai_result.get("tips", []),
    }